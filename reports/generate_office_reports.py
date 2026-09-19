#!/usr/bin/env python3
"""Generate separated office-style Excel and Word reports for NARARYA BUSINESS.
Sensitive identity documents are intentionally excluded from these reports.
"""
from __future__ import annotations
from pathlib import Path
from datetime import datetime, timezone
from typing import Iterable, Mapping

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("Install dependency: pip install openpyxl python-docx") from exc

try:
    from docx import Document
except ImportError as exc:
    raise SystemExit("Install dependency: pip install python-docx") from exc

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "generated"
OUT.mkdir(parents=True, exist_ok=True)
FOOTER = "PT NEXOVONARSACORPORATION - All Right Reserved"

DATASETS = {
    "sales": ("NARARYA_SALES.xlsx", "NARARYA_SALES.docx",
              ["sale_id","date","brand","product","customer_ref","amount","payment_status"]),
    "orders": ("NARARYA_ORDERS.xlsx", "NARARYA_ORDERS.docx",
               ["order_id","created_at","brand","product","customer_ref","amount","status"]),
    "finance": ("NARARYA_FINANCE.xlsx", "NARARYA_FINANCE.docx",
                ["transaction_id","date","kind","brand","category","description","amount","method","order_id","status"]),
    "members": ("NARARYA_MEMBERS.xlsx", "NARARYA_MEMBERS.docx",
                ["member_id","status","name","whatsapp","gmail","instagram","tiktok","po_bussid","po_ets2","consent_at","secure_evidence_ref"]),
    "catalog": ("NARARYA_CATALOG.xlsx", "NARARYA_CATALOG.docx",
                ["retailer_id","brand","name","category","price","currency","availability","description"]),
}

def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat()

def make_workbook(title: str, fields: list[str], rows: Iterable[Mapping[str, object]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.freeze_panes = "A2"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    for idx, field in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(field)+2, 14), 42)
    return wb

def make_word(title: str, fields: list[str], rows: list[Mapping[str, object]]) -> Document:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Generated: " + now_utc())
    doc.add_paragraph("Laporan operasional terpisah. Dokumen identitas sensitif mentah tidak dimasukkan.")
    if rows:
        table = doc.add_table(rows=1, cols=len(fields))
        for i, field in enumerate(fields):
            table.rows[0].cells[i].text = field
        for row in rows:
            cells = table.add_row().cells
            for i, field in enumerate(fields):
                cells[i].text = str(row.get(field, ""))
    else:
        doc.add_paragraph("Belum ada data.")
    doc.add_paragraph(FOOTER)
    return doc

def generate(dataset: str, rows: Iterable[Mapping[str, object]] = ()) -> tuple[Path, Path]:
    if dataset not in DATASETS:
        raise ValueError("Unknown dataset: " + dataset)
    xlsx_name, docx_name, fields = DATASETS[dataset]
    safe_rows = list(rows)
    xlsx_path = OUT / xlsx_name
    make_workbook(dataset.upper(), fields, safe_rows).save(xlsx_path)
    docx_path = OUT / docx_name
    make_word(dataset.upper(), fields, safe_rows).save(docx_path)
    return xlsx_path, docx_path

def generate_all(data: Mapping[str, Iterable[Mapping[str, object]]] | None = None) -> list[tuple[Path, Path]]:
    source = data or {key: [] for key in DATASETS}
    return [generate(name, source.get(name, [])) for name in DATASETS]

if __name__ == "__main__":
    for xlsx, docx in generate_all():
        print("READY", xlsx)
        print("READY", docx)
